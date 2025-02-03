import requests
from pystyle import Colors, Write
from phonenumbers import geocoder, carrier
import phonenumbers
import os
import socket
import ssl
from concurrent.futures import ThreadPoolExecutor, as_completed
import dns.resolver
from dns import reversename
from email_validator import validate_email, EmailNotValidError
from urllib.parse import quote
import secrets
import json
from bs4 import BeautifulSoup
import re
from email.parser import Parser
import whois
from tqdm import tqdm
from datetime import datetime
import openai

import magic
import stat
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS
import PyPDF2

import openpyxl
import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx import Presentation

from mutagen.easyid3 import EasyID3
from mutagen.mp3 import MP3
from mutagen.mp4 import MP4
from mutagen.id3 import ID3
from mutagen.flac import FLAC
import wave
from mutagen.oggvorbis import OggVorbis
from tinytag import TinyTag

default_color = Colors.red
API_KEY = "INSERT GOOGLE SEARCH API KEY HERE.."
CX = "INSERT GOOGLE SEARCH CX KEY HERE"
CLIENT_ID = "INSERT GOOGLE SEARCH CLIENT ID KEY HERE"
HIBP_API_KEY = "INSERT HAVE I BEEN PWNED API KEY HERE"

_global_session = requests.Session()
default_color = Colors.light_red
requests.get = _global_session.get

import multiprocessing
MAX_WORKERS = min(32, (multiprocessing.cpu_count() or 1) * 5)

def subdomain_enumeration(domain):
    import requests
    from datetime import datetime
    url = f"https://crt.sh/?q=%.{domain}&output=json"
    Write.Print(f"\n[!] Subdomain enumeration for: {domain}\n", Colors.white, interval=0)
    try:
        resp = requests.get(url, timeout=10)
        if resp.status_code == 200:
            try:
                data = resp.json()
            except json.JSONDecodeError:
                msg = "[!] > Error: crt.sh returned non-JSON or empty data.\n"
                Write.Print(msg, Colors.red, interval=0)
                return
            found_subs = set()
            for entry in data:
                if 'name_value' in entry:
                    for subd in entry['name_value'].split('\n'):
                        subd_strip = subd.strip()
                        if subd_strip and subd_strip != domain:
                            found_subs.add(subd_strip)
                elif 'common_name' in entry:
                    c = entry['common_name'].strip()
                    if c and c != domain:
                        found_subs.add(c)
            if found_subs:
                out_text = f"\n[+] Found {len(found_subs)} subdomains for {domain}:\n"
                for s in sorted(found_subs):
                    out_text += f"    {s}\n"
                Write.Print(out_text, Colors.green, interval=0)
                print()
                print("[?] Would you like to save this output to a log file? (Y/N): ", end="")
                choice = input().strip().upper()
                if choice == 'Y':
                    stamp = datetime.now().strftime("[%Y-%m-%d %H:%M:%S] ")
                    with open("clatscope_log.txt", "a", encoding="utf-8") as f:
                        f.write(stamp + out_text + "\n")
                    Write.Print("[!] > Subdomains saved to clatscope_log.txt\n", Colors.white, interval=0)
            else:
                Write.Print("[!] > No subdomains found.\n", Colors.red, interval=0)
        else:
            err = f"[!] > HTTP {resp.status_code} from crt.sh\n"
            Write.Print(err, Colors.red, interval=0)
    except Exception as exc:
        Write.Print(f"[!] > Subdomain enumeration error: {exc}\n", Colors.red, interval=0)

def validate_domain_input(domain):
    if not domain or len(domain) > 253 or ".." in domain:
        return False
    pattern = r"^[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$"
    return bool(re.match(pattern, domain))

def log_option(output_text):
    print()
    print("[?] Would you like to save this output to a log file? (Y/N): ", end="")
    choice = input().strip().upper()
    if choice == 'Y':
        stamp = datetime.now().strftime("[%Y-%m-%d %H:%M:%S] ")
        with open("clatscope_log.txt", "a", encoding="utf-8") as log_file:
            log_file.write(f"{stamp}{output_text}\n\n")
        Write.Print("[!] > Output has been saved to clatscope_log.txt\n", default_color, interval=0)

def clear():
    os.system('cls' if os.name == 'nt' else 'clear')

def restart():
    Write.Input("\nPress Enter to return to the main menu...", default_color, interval=0)
    clear()

def get_ip_details(ip):
    try:
        response = requests.get(f"https://ipinfo.io/{ip}/json", timeout=10)
        response.raise_for_status()
        return response.json()
    except:
        return None

def person_search(first_name, last_name, city):
    query = f"{first_name} {last_name} {city}"
    payload_person_search = {
        "model": "llama-3.1-sonar-huge-128k-online",
        "messages": [
            {
                "role": "system",
                "content": (
                    "When searching for information about [PERSON NAME], please provide a comprehensive analysis that includes:"
                    "their full name including any variations or aliases; date and place of birth (if known); current location or"
                    "place of death if deceased; educational background including institutions and years of attendance; professional"
                    "history with dates and notable achievements; any significant public roles or positions held; major life events"
                    "or controversies; family connections and relationships relevant to their public life; and their current status"
                    "or most recent known activities. For any claims made, include specific citations using [Source X] notation"
                    "within the text, where X corresponds to the numbered source in the reference list. Each piece of information"
                    "should be attributed to at least one credible source, with preference given to primary sources, official records,"
                    "reputable news organizations, and peer-reviewed academic works where applicable. Avoid speculation beyond verifiable"
                    "data. When researching [PERSON NAME], first verify the specific individual by their distinguishing characteristics"
                    "(occupation, time period, location, or notable achievements). If the person has appeared in a news article or public interview, discuss the details of it."
                    "If multiple people share similar names, acknowledge their existence at the beginning of your response like this: Note: There are other notable individuals named"
                    "[Similar Name], including [Brief one-line identifier for each]. This analysis focuses on [Target Person] who is known for [Key Identifier]."
                    "Then proceed with the detailed analysis of only the target individual, including their background, achievements, and current status."
                    "If you cannot confidently distinguish between similarly named individuals based on the available context, state this uncertainty clearly"
                    "and list the potential matches with their key identifiers, requesting additional details to ensure accurate identification. All information"
                    "should be properly cited using numbered references, and only include verified information about the specific target individual. At the end of the analysis,"
                    "provide a numbered list of all sources cited, including full bibliographic information (author, title, publication, date, URL if applicable) in Chicago style format."
                    "If any critical information is missing orunverifiable, explicitly note these gaps in the analysis. Include information about their current job, employment status, and other relevant professional information."
                )
             },
            {
                "role": "user",
                "content": f"Provide detailed background or publicly known information about: {query}"
            }
        ],
        "max_tokens": 50000,
        "temperature": 1.0
    }
    PERPLEXITY_API_URL = "https://api.perplexity.ai/chat/completions"
    PERPLEXITY_API_KEY = "INSERT PERPLEXITY API KEY HERE"
    perplexity_headers = {
        "Authorization": f"Bearer {PERPLEXITY_API_KEY}",
        "Content-Type": "application/json",
    }
    results_text = ""
    try:
        response = requests.post(PERPLEXITY_API_URL, headers=perplexity_headers, json=payload_person_search)
        if response.status_code == 200:
            data = response.json()
            info_content = data["choices"][0]["message"]["content"]
            results_text = (
                f"\nPERSON SEARCH RESULTS\n"
                f"=====================\n\n"
                f"NAME:\n{first_name} {last_name}\n\n"
                f"LOCATION:\n{city}\n\n"
                f"PUBLIC INFORMATION:\n{info_content}\n"
            )
        else:
            results_text = f"[!] > Error from Perplexity: HTTP {response.status_code}\n{response.text}\n"
    except Exception as e:
        results_text = f"[!] > Error: {str(e)}\n"
    clear()
    Write.Print(results_text, Colors.white, interval=0)
    log_option(results_text)
    restart()

def fetch_page_text(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36)"
    }
    try:
        resp = requests.get(url, headers=headers, timeout=20)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        for tag_name in ["header", "footer", "nav", "aside", "script", "style", "noscript", "form"]:
            for t in soup.find_all(tag_name):
                t.decompose()
        text = soup.get_text(separator=' ')
        text = ' '.join(text.split())
        return text if text else "No meaningful content found."
    except Exception:
        return "Could not retrieve or parse the webpage content."

def ip_info(ip):
    url = f"https://ipinfo.io/{ip}/json"
    try:
        response = requests.get(url)
        response.raise_for_status()
        data = response.json()
        loc = data.get('loc', 'None')
        maps_link = f"https://www.google.com/maps?q={loc}" if loc != 'None' else 'None'
        ip_details = f"""
╭─{' '*78}─╮
|{' '*34} IP Details {' '*34}|
|{'='*80}|
| [+] > IP Address         || {data.get('ip', 'None'):<51}|
| [+] > City               || {data.get('city', 'None'):<51}|
| [+] > Region             || {data.get('region', 'None'):<51}|
| [+] > Country            || {data.get('country', 'None'):<51}|
| [+] > Postal/ZIP Code    || {data.get('postal', 'None'):<51}|
| [+] > ISP                || {data.get('org', 'None'):<51}|
| [+] > Coordinates        || {loc:<51}|
| [+] > Timezone           || {data.get('timezone', 'None'):<51}|
| [+] > Location           || {maps_link:<51}|
╰─{' '*24}─╯╰─{' '*50}─╯
"""
        Write.Print(ip_details, Colors.white, interval=0)
        log_option(ip_details)
    except:
        clear()
        Write.Print("\n[!] > Error retrieving IP address info.", default_color, interval=0)
    restart()

def fetch_social_urls(urls, title):
    def check_url(url):
        try:
            response = requests.get(url, timeout=10)
            status_code = response.status_code
            if status_code == 200:
                return f"[+] > {url:<50}|| Found"
            elif status_code == 404:
                return f"[-] > {url:<50}|| Not found"
            else:
                return f"[-] > {url:<50}|| Error: {status_code}"
        except requests.exceptions.Timeout:
            return f"[-] > {url:<50}|| Timeout"
        except requests.exceptions.ConnectionError:
            return f"[-] > {url:<50}|| Connection error"
        except requests.exceptions.RequestException:
            return f"[-] > {url:<50}|| Request error"
        except Exception:
            return f"[-] > {url:<50}|| Unexpected error"
    result_str = f"""
╭─{' '*78}─╮
|{' '*27}{title}{' '*27}|
|{'='*80}|
"""
    with ThreadPoolExecutor() as executor:
        executor._max_workers = MAX_WORKERS
        results = list(executor.map(check_url, urls))
    for result in results:
        result_str += f"| {result:<78} |\n"
    result_str += f"╰─{' '*78}─╯"
    return result_str

def deep_account_search(nickname):
    sites = [
        "https://youtube.com/@{target}",
        "https://facebook.com/{target}",
        "https://wikipedia.org/wiki/User:{target}",
        "https://instagram.com/{target}",
        "https://reddit.com/user/{target}",
        "https://medium.com/@{target}",
        "https://www.quora.com/profile/{target}",
        "https://bing.com/{target}",
        "https://x.com/{target}",
        "https://yandex.ru/{target}",
        "https://whatsapp.com/{target}",
        "https://yahoo.com/{target}",
        "https://amazon.com/{target}",
        "https://duckduckgo.com/{target}",
        "https://yahoo.co.jp/{target}",
        "https://tiktok.com/@{target}",
        "https://msn.com/{target}",
        "https://netflix.com/{target}",
        "https://weather.com/{target}",
        "https://live.com/{target}",
        "https://naver.com/{target}",
        "https://microsoft.com/{target}",
        "https://twitch.tv/{target}",
        "https://office.com/{target}",
        "https://vk.com/{target}",
        "https://pinterest.com/{target}",
        "https://discord.com/{target}",
        "https://aliexpress.com/{target}",
        "https://github.com/{target}",
        "https://adobe.com/{target}",
        "https://rakuten.co.jp/{target}",
        "https://ikea.com/{target}",
        "https://bbc.co.uk/{target}",
        "https://amazon.co.jp/{target}",
        "https://speedtest.net/{target}",
        "https://samsung.com/{target}",
        "https://healthline.com/{target}",
        "https://medlineplus.gov/{target}",
        "https://roblox.com/users/{target}/profile",
        "https://cookpad.com/{target}",
        "https://indiatimes.com/{target}",
        "https://mercadolivre.com.br/{target}",
        "https://britannica.com/{target}",
        "https://merriam-webster.com/{target}",
        "https://hurriyet.com.tr/{target}",
        "https://steamcommunity.com/id/{target}",
        "https://booking.com/{target}",
        "https://support.google.com/{target}",
        "https://bbc.com/{target}",
        "https://playstation.com/{target}",
        "https://ebay.com/usr/{target}",
        "https://poki.com/{target}",
        "https://walmart.com/{target}",
        "https://medicalnewstoday.com/{target}",
        "https://gov.uk/{target}",
        "https://nhs.uk/{target}",
        "https://detik.com/{target}",
        "https://cricbuzz.com/{target}",
        "https://nih.gov/{target}",
        "https://uol.com.br/{target}",
        "https://ilovepdf.com/{target}",
        "https://clevelandclinic.org/{target}",
        "https://cnn.com/{target}",
        "https://globo.com/{target}",
        "https://nytimes.com/{target}",
        "https://taboola.com/{target}",
        "https://pornhub.com/users/{target}",
        "https://redtube.com/users/{target}",
        "https://xnxx.com/profiles/{target}",
        "https://brazzers.com/profile/{target}",
        "https://xhamster.com/users/{target}",
        "https://onlyfans.com/{target}",
        "https://xvideos.es/profiles/{target}",
        "https://xvideos.com/profiles/{target}",
        "https://chaturbate.com/{target}",
        "https://redgifs.com/users/{target}",
        "https://tinder.com/{target}",
        "https://pof.com/{target}",
        "https://match.com/{target}",
        "https://eharmony.com/{target}",
        "https://bumble.com/{target}",
        "https://okcupid.com/{target}",
        "https://Badoo.com/{target}",
        "https://dating.com/{target}",
        "https://trello.com/{target}",
        "https://mapquest.com/{target}",
        "https://zoom.com/{target}",
        "https://apple.com/{target}",
        "https://dropbox.com/{target}",
        "https://weibo.com/{target}",
        "https://wordpress.com/{target}",
        "https://cloudflare.com/{target}",
        "https://salesforce.com/{target}",
        "https://fandom.com/{target}",
        "https://paypal.com/{target}",
        "https://soundcloud.com/{target}",
        "https://forbes.com/{target}",
        "https://theguardian.com/{target}",
        "https://hulu.com/{target}",
        "https://stackoverflow.com/users/{target}",
        "https://businessinsider.com/{target}",
        "https://huffpost.com/{target}",
        "https://booking.com/{target}",
        "https://pastebin.com/u/{target}",
        "https://producthunt.com/@{target}",
        "https://pypi.org/user/{target}",
        "https://slideshare.com/{target}",
        "https://strava.com/athletes/{target}",
        "https://tldrlegal.com/{target}",
        "https://t.me/{target}",
        "https://last.fm/user{target}",
        "https://data.typeracer.com/pit/profile?user={target}",
        "https://tryhackme.com/p/{target}",
        "https://trakt.tv/users/{target}",
        "https://scratch.mit.edu/users/{target}",
        "https://replit.com?{target}",
        "https://hackaday.io/{target}",
        "https://freesound.org/people/{target}",
        "https://hub.docker.com/u/{target}",
        "https://disqus.com/{target}",
        "https://www.codecademy.com/profiles/{target}",
        "https://www.chess.com/member/{target}",
        "https://bitbucket.org/{target}",
        "https://www.twitch.tv?{target}",
        "https://wikia.com/wiki/User:{target}",
        "https://steamcommunity.com/groups{target}",
        "https://keybase.io?{target}",
        "http://en.gravatar.com/{target}",
        "https://vk.com/{target}",
        "https://deviantart.com/{target}",
        "https://www.behance.net/{target}",
        "https://vimeo.com/{target}",
        "https://www.youporn.com/user/{target}",
        "https://profiles.wordpress.org/{target}",
        "https://tryhackme.com/p/{target}",
        "https://www.scribd.com/{target}",
        "https://myspace.com/{target}",
        "https://genius.com/{target}",
        "https://genius.com/artists/{target}",
        "https://www.flickr.com/people/{target}",
        "https://www.fandom.com/u/{target}",
        "https://www.chess.com/member/{target}",
        "https://buzzfeed.com/{target}",
        "https://www.buymeacoffee.com/{target}",
        "https://about.me/{target}",
        "https://discussions.apple.com/profile/{target}",
        "https://giphy.com/{target}",
        "https://scholar.harvard.edu/{target}",
        "https://www.instructables.com/member/{target}",
        "http://www.wikidot.com/user:info/{target}",
        "https://www.youporn.com/user/{target}",
        "https://profiles.wordpress.org/{target}",
        "https://vimeo.com/{target}",
        "https://tryhackme.com/p/{target}",
        "https://www.scribd.com/{target}",
        "https://myspace.com/{target}",
        "https://hackerone.com/{target}",
        "https://www.hackerearth.com/@{target}",
        "https://genius.com/{target}",
        "https://genius.com/artists/{target}",
        "https://www.flickr.com/people/{target}",
        "https://www.fandom.com/u/{target}",
        "https://www.chess.com/member/{target}",
        "https://buzzfeed.com/{target}",
        "https://about.me/{target}",
        "https://discussions.apple.com/profile/{target}",
        "https://archive.org/details/@{target}",
        "https://giphy.com/{target}",
        "https://scholar.harvard.edu/{target}",
        "https://www.instructables.com/member/{target}",
        "http://www.wikidot.com/user:info/{target}",
        "https://dribbble.com/{target}",
        "https://500px.com/{target}",
        "https://disqus.com/by/{target}",
        "https://steamcommunity.com/id{target}",
        "https://www.codecademy.com/profiles/{target}",
        "https://www.khanacademy.org/profile/{target}",
        "https://bitbucket.org/{target}/",
        "https://sourceforge.net/u/{target}/profile/",
        "https://www.gog.com/u/{target}",
        "https://www.dailymotion.com/{target}",
        "https://myanimelist.net/profile/{target}",
        "https://archiveofourown.org/users/{target}",
        "https://www.slideshare.net/{target}",
        "https://letterboxd.com/{target}/",
        "https://runkeeper.com/{target}/profile",
        "https://www.mapmyrun.com/profile/{target}",
        "https://patreon.com/{target}",
        "https://www.producthunt.com/@{target}",
        "https://about.me/{target}",
        "https://www.scribd.com/{target}",
        "https://www.mixcloud.com/{target}/",
        "https://keybase.io/{target}",
        "https://www.houzz.com/user/{target}",
        "https://www.couchsurfing.com/people{target}",
        "https://mix.com/{target}",
        "https://medium.com/@{target}",
        "https://getpocket.com/@{target}",
        "https://www.instructables.com/member/{target}/",
        "https://colorlib.com/author{target}/",
        "https://en.gravatar.com/{target}",
        "https://www.kaggle.com/{target}",
        "https://replit.com/@{target}",
        "https://scratch.mit.edu/users/{target}",
        "https://www.alik.cz/u/{target}",
        "https://www.munzee.com/m/{target}",
        "https://boardgamegeek.com/user/{target}",
        "https://www.youporn.com/user/{target}",
        "https://account.xbox.com/en-us/profile?gamertag={target}",
        "https://profiles.wordpress.org/{target}",
        "https://vimeo.com/{target}",
        "https://tryhackme.com/p/{target}",
        "https://torrentgalaxy.to/profile/{target}",
        "https://www.scribd.com/{target}",
        "https://www.patreon.com/{target}",
        "https://myspace.com/{target}",
        "https://genius.com/{target}",
        "https://genius.com/artists/{target}",
        "https://www.flickr.com/people/{target}",
        "https://www.fandom.com/u/{target}",
        "https://www.chess.com/member/{target}",
        "https://buzzfeed.com/{target}",
        "https://www.buymeacoffee.com/{target}",
        "https://about.me/{target}",
        "https://discussions.apple.com/profile/{target}",
        "https://archive.org/details/@{target}",
        "https://giphy.com/{target}",
        "https://scholar.harvard.edu/{target}",
        "https://www.instructables.com/member/{target}",
        "http://www.wikidot.com/user:info/{target}",
        "https://erome.com/{target}",
        "https://www.alik.cz/u/{target}",
        "https://rblx.trade/p/{target}",
        "https://www.paypal.com/paypalme/{target}",
        "https://hackaday.io/{target}",
        "https://connect.garmin.com/modern/profile/{target}"
    ]
    urls = [site_format.format(target=nickname) for site_format in sites]
    search_results = fetch_social_urls(urls, "Deep Account Search")
    Write.Print(search_results, Colors.white, interval=0)
    log_option(search_results)
    restart()

def phone_info(phone_number):
    try:
        parsed_number = phonenumbers.parse(phone_number)
        country = geocoder.country_name_for_number(parsed_number, "en")
        region = geocoder.description_for_number(parsed_number, "en")
        operator = carrier.name_for_number(parsed_number, "en")
        valid = phonenumbers.is_valid_number(parsed_number)
        validity = "Valid" if valid else "Invalid"
        phonetext = f"""
\n
╭─{' '*50}─╮
|{' '*17}Phone number info{' '*18}|
|{'='*52}|
| [+] > Number   || {phone_number:<33}|
| [+] > Country  || {country:<33}     |
| [+] > Region   || {region:<33}      |
| [+] > Operator || {operator:<33}    |
| [+] > Validity || {validity:<33}    |
╰─{' '*15}─╯╰─{' '*31}─╯\n"""
        Write.Print(phonetext, Colors.white, interval=0)
        log_option(phonetext)
    except phonenumbers.phonenumberutil.NumberParseException:
        clear()
        Write.Print(f"\n[!] > Error: invalid phone number format (+1-000-000-0000)", default_color, interval=0)
    restart()

def dns_lookup(domain):
    record_types = ['A', 'CNAME', 'MX', 'NS']
    result_output = f"""
╭─{' '*78}─╮
|{' '*33} DNS Lookup {' '*33}|
|{'='*80}|
"""
    for rtype in record_types:
        result_output += f"| [+] > {rtype} Records: {' '*62}|\n"
        try:
            answers = dns.resolver.resolve(domain, rtype)
            for ans in answers:
                if rtype == 'MX':
                    result_output += f"|    {ans.preference:<4} {ans.exchange:<70}|\n"
                else:
                    result_output += f"|    {str(ans):<76}|\n"
        except dns.resolver.NoAnswer:
            result_output += "|    No records found.\n"
        except dns.resolver.NXDOMAIN:
            result_output += "|    Domain does not exist.\n"
        except Exception:
            result_output += "|    Error retrieving records.\n"
        result_output += f"|{'='*80}|\n"
    result_output += f"╰─{' '*78}─╯"
    Write.Print(result_output, Colors.white, interval=0)
    log_option(result_output)
    restart()

def email_lookup(email_address):
    try:
        v = validate_email(email_address)
        email_domain = v.domain
    except EmailNotValidError as e:
        Write.Print(f"[!] > Invalid email address format: {str(e)}", default_color, interval=0)
        restart()
        return
    mx_records = []
    try:
        answers = dns.resolver.resolve(email_domain, 'MX')
        for rdata in answers:
            mx_records.append(str(rdata.exchange))
    except:
        mx_records = []
    validity = "Mx Found (Might be valid)" if mx_records else "No MX found (Might be invalid)"
    email_text = f"""
╭─{' '*78}─╮
|{' '*34}Email Info{' '*34}|
|{'='*80}|
| [+] > Email:        || {email_address:<52}|
| [+] > Domain:       || {email_domain:<52}|
| [+] > MX Records:   || {", ".join(mx_records) if mx_records else "None":<52}|
| [+] > Validity:     || {validity:<52}|
╰─{' '*23}─╯╰─{' '*51}─╯
"""
    Write.Print(email_text, Colors.white, interval=0)
    log_option(email_text)
    restart()

def reverse_dns(ip):
    try:
        rev_name = reversename.from_address(ip)
        answers = dns.resolver.resolve(rev_name, "PTR")
        ptr_record = str(answers[0]).strip('.')
    except:
        ptr_record = "No PTR record found"
    rdns_text = f"""
╭─{' '*78}─╮
|{' '*33}Reverse DNS Lookup{' '*33}|
|{'='*80}|
| [+] > IP:     || {ip:<60}|
| [+] > Host:   || {ptr_record:<60}|
╰─{' '*23}─╯╰─{' '*51}─╯
"""
    Write.Print(rdns_text, Colors.white, interval=0)
    log_option(rdns_text)
    restart()

def analyze_email_header(raw_headers):
    parser = Parser()
    msg = parser.parsestr(raw_headers)
    from_ = msg.get("From", "")
    to_ = msg.get("To", "")
    subject_ = msg.get("Subject", "")
    date_ = msg.get("Date", "")
    received_lines = msg.get_all("Received", [])
    found_ips = []
    if received_lines:
        for line in received_lines:
            potential_ips = re.findall(r'\b\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}\b', line)
            for ip in potential_ips:
                if ip not in found_ips:
                    found_ips.append(ip)
    header_text = f"""
╭─{' '*78}─╮
|{' '*31}Email Header Analysis{' '*31}|
|{'='*80}|
| [+] > From:      || {from_:<55}|
| [+] > To:        || {to_:<55}|
| [+] > Subject:   || {subject_:<55}|
| [+] > Date:      || {date_:<55}|
|{'-'*80}|
"""
    if found_ips:
        header_text += "| [+] > Received Path (IPs found):\n"
        for ip in found_ips:
            header_text += f"|    {ip:<76}|\n"
    else:
        header_text += "| [+] > No IPs found in Received headers.\n"
    header_text += f"╰─{' '*78}─╯"
    Write.Print(header_text, Colors.white, interval=0)
    if found_ips:
        ip_details_header = f"""
╭─{' '*78}─╮
|{' '*30}IP Geolocation Details{' '*30}|
|{'='*80}|
"""
        ip_details_summary = ""
        for ip in found_ips:
            data = get_ip_details(ip)
            if data is not None:
                loc = data.get('loc', 'None')
                ip_details_summary += f"| IP: {ip:<14}|| City: {data.get('city','N/A'):<15} Region: {data.get('region','N/A'):<15} Country: {data.get('country','N/A'):<4}|\n"
                ip_details_summary += f"|    Org: {data.get('org','N/A'):<63}|\n"
                ip_details_summary += f"|    Loc: {loc:<63}|\n"
                ip_details_summary += "|" + "-"*78 + "|\n"
            else:
                ip_details_summary += f"| IP: {ip:<14}|| [!] Could not retrieve details.\n"
                ip_details_summary += "|" + "-"*78 + "|\n"
        ip_details_footer = f"╰─{' '*78}─╯"
        ip_details_full = ip_details_header + ip_details_summary + ip_details_footer
        Write.Print(ip_details_full, Colors.white, interval=0)
        header_text += "\n" + ip_details_full
    spf_result, dkim_result, dmarc_result = None, None, None
    spf_domain, dkim_domain = None, None
    auth_results = msg.get_all("Authentication-Results", [])
    from_domain = ""
    if "@" in from_:
        from_domain = from_.split("@")[-1].strip(">").strip()
    if auth_results:
        for entry in auth_results:
            spf_match = re.search(r'spf=(pass|fail|softfail|neutral)', entry, re.IGNORECASE)
            if spf_match:
                spf_result = spf_match.group(1)
            spf_domain_match = re.search(r'envelope-from=([^;\s]+)', entry, re.IGNORECASE)
            if spf_domain_match:
                spf_domain = spf_domain_match.group(1)
            dkim_match = re.search(r'dkim=(pass|fail|none|neutral)', entry, re.IGNORECASE)
            if dkim_match:
                dkim_result = dkim_match.group(1)
            dkim_domain_match = re.search(r'd=([^;\s]+)', entry, re.IGNORECASE)
            if dkim_domain_match:
                dkim_domain = dkim_domain_match.group(1)
            dmarc_match = re.search(r'dmarc=(pass|fail|none)', entry, re.IGNORECASE)
            if dmarc_match:
                dmarc_result = dmarc_match.group(1)
    spf_align = False
    dkim_align = False
    if from_domain and spf_domain:
        spf_align = from_domain.lower() == spf_domain.lower()
    if from_domain and dkim_domain:
        dkim_align = from_domain.lower() == dkim_domain.lower()
    alignment_text = f"""
╭─{' '*78}─╮
|{' '*30}SPF / DKIM / DMARC Checks{' '*29}|
|{'='*80}|
| [+] > SPF  Result:   {spf_result if spf_result else 'Not found':<20}   Domain: {spf_domain if spf_domain else 'N/A':<20} Aligned: {spf_align}|
| [+] > DKIM Result:   {dkim_result if dkim_result else 'Not found':<20} Domain: {dkim_domain if dkim_domain else 'N/A':<20} Aligned: {dkim_align}|
| [+] > DMARC Result:  {dmarc_result if dmarc_result else 'Not found':<20}|
╰─{' '*78}─╯
"""
    Write.Print(alignment_text, Colors.white, interval=0)
    full_output = header_text + "\n" + alignment_text
    log_option(full_output)
    restart()

def haveibeenpwned_check(email):
    headers = {
        "hibp-api-key": HIBP_API_KEY,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36)"
    }
    url = f"https://haveibeenpwned.com/api/v3/breachedaccount/{email}?truncateResponse=false"
    try:
        resp = requests.get(url, headers=headers, timeout=10)
        if resp.status_code == 200:
            breaches = resp.json()
            clear()
            results_text = f"""
╭─{' '*78}─╮
|{' '*30}Have I Been Pwned?{' '*30}|
|{'='*80}|
| [!] > Bad news! Your email was found in {len(breaches)} breach(es)                          |
|{'-'*80}|
"""
            for index, breach in enumerate(breaches, start=1):
                breach_name = breach.get('Name', 'Unknown')
                domain = breach.get('Domain', 'Unknown')
                breach_date = breach.get('BreachDate', 'Unknown')
                added_date = breach.get('AddedDate', 'Unknown')
                pwn_count = breach.get('PwnCount', 'Unknown')
                data_classes = ", ".join(breach.get('DataClasses', []))
                results_text += f"| Breach #{index}: {breach_name:<66}|\n"
                results_text += f"|    Domain: {domain:<71}|\n"
                results_text += f"|    Breach Date: {breach_date:<65}|\n"
                results_text += f"|    Added Date:  {added_date:<65}|\n"
                results_text += f"|    PwnCount:    {pwn_count:<65}|\n"
                results_text += f"|    Data Types:  {data_classes:<65}|\n"
                results_text += f"|{'='*80}|\n"
            results_text += f"╰─{' '*78}─╯"
            Write.Print(results_text, Colors.white, interval=0)
            log_option(results_text)
        elif resp.status_code == 404:
            clear()
            msg = f"""
╭─{' '*78}─╮
|{' '*30}Have I Been Pwned?{' '*30}|
|{'='*80}|
| [!] > Good news! No breaches found for: {email:<48}|
╰─{' '*78}─╯
"""
            Write.Print(msg, Colors.white, interval=0)
            log_option(msg)
        else:
            clear()
            error_msg = f"""
[!] > An error occurred: HTTP {resp.status_code}
Response: {resp.text}
"""
            Write.Print(error_msg, Colors.red, interval=0)
            log_option(error_msg)
    except requests.exceptions.Timeout:
        clear()
        Write.Print("[!] > Request timed out when contacting Have I Been Pwned.", default_color, interval=0)
    except Exception as e:
        clear()
        Write.Print(f"[!] > An error occurred: {str(e)}", default_color, interval=0)
    restart()

def change_color():
    global default_color
    clear()
    color_menu = """
╭─    ─╮╭─                     ─╮
|  №   ||         Color         |
|======||=======================|
| [1]  || Red                   |
| [2]  || Blue                  |
| [3]  || Green                 |
| [4]  || Yellow                |
| [5]  || Cyan                  |
| [6]  || White                 |
|------||-----------------------|
| [0]  || Back to settings menu |
╰─    ─╯╰─                     ─╯
"""
    Write.Print(color_menu, Colors.white, interval=0)
    color_choice = Write.Input("\n\n[?] >  ", default_color, interval=0).strip()
    color_map = {
        "1": Colors.red,
        "2": Colors.blue,
        "3": Colors.green,
        "4": Colors.yellow,
        "5": Colors.cyan,
        "6": Colors.white
    }
    if color_choice in color_map:
        default_color = color_map[color_choice]
        clear()
        Write.Print("[!] > Colour has been changed.\n", default_color, interval=0)
    elif color_choice == "0":
        settings()
    else:
        clear()
        Write.Print("[!] > Invalid choice.\n", Colors.red, interval=0)
    restart()

def whois_lookup(domain):
    try:
        w = whois.whois(domain)
        clear()
        domain_name = w.domain_name if w.domain_name else "N/A"
        registrar = w.registrar if w.registrar else "N/A"
        creation_date = w.creation_date if w.creation_date else "N/A"
        expiration_date = w.expiration_date if w.expiration_date else "N/A"
        updated_date = w.updated_date if w.updated_date else "N/A"
        name_servers = ", ".join(w.name_servers) if w.name_servers else "N/A"
        status = ", ".join(w.status) if w.status else "N/A"
        whois_text = f"""
╭─{' '*78}─╮
|{' '*34}WHOIS Lookup{' '*34}|
|{'='*80}|
| [+] > Domain Name:       || {str(domain_name):<52}|
| [+] > Registrar:         || {str(registrar):<52}|
| [+] > Creation Date:     || {str(creation_date):<52}|
| [+] > Expiration Date:   || {str(expiration_date):<52}|
| [+] > Updated Date:      || {str(updated_date):<52}|
| [+] > Name Servers:      || {name_servers:<52}|
| [+] > Status:            || {status:<52}|
╰─{' '*23}─╯╰─{' '*51}─╯
"""
        Write.Print(whois_text, Colors.white, interval=0)
        log_option(whois_text)
    except Exception as e:
        clear()
        Write.Print(f"[!] > WHOIS lookup error: {str(e)}", default_color, interval=0)
    restart()

def check_password_strength(password):
    txt_file_path = os.path.join(os.path.dirname(__file__), "passwords.txt")
    if os.path.isfile(txt_file_path):
        try:
            with open(txt_file_path, "r", encoding="utf-8") as f:
                common_words = f.read().splitlines()
            for word in common_words:
                if word and word in password:
                    return "Weak password (may contain common phrase, term, word, sequence, etc, DO NOT use this password)"
        except Exception:
            pass
    score = 0
    if len(password) >= 8:
        score += 1
    if len(password) >= 12:
        score += 1
    if re.search(r'[A-Z]', password):
        score += 1
    if re.search(r'[a-z]', password):
        score += 1
    if re.search(r'\d', password):
        score += 1
    if re.search(r'[^a-zA-Z0-9]', password):
        score += 1
    if score <= 2:
        return "Weak password (may contain common phrase, term, word, sequence, etc, DO NOT use this password)"
    elif 3 <= score <= 4:
        return "Moderate password (room for improvement)"
    else:
        return "Strong password (suitable for high security apps / credentials)"

def password_strength_tool():
    clear()
    Write.Print("[!] > Enter password to evaluate strength:\n", default_color, interval=0)
    password = Write.Input("[?] >  ", default_color, interval=0)
    if not password:
        clear()
        Write.Print("[!] > Password cannot be empty. Please enter the password.\n", default_color, interval=0)
        restart()
        return
    strength = check_password_strength(password)
    clear()
    output_text = f"Password Strength: {strength}\n"
    Write.Print(output_text, Colors.white, interval=0)
    log_option(output_text)
    restart()

def fetch_wmn_data():
    try:
        response = requests.get("https://raw.githubusercontent.com/WebBreacher/WhatsMyName/main/wmn-data.json", timeout=10)
        response.raise_for_status()
        return response.json()
    except:
        Write.Print("[!] > Failed to fetch data from WhatsMyName repository.\n", Colors.red, interval=0)
        return None

def check_site(site, username, headers):
    site_name = site["name"]
    uri_check = site["uri_check"].format(account=username)
    try:
        res = requests.get(uri_check, headers=headers, timeout=10)
        estring_pos = site["e_string"] in res.text
        estring_neg = site["m_string"] in res.text
        if res.status_code == site["e_code"] and estring_pos and not estring_neg:
            return site_name, uri_check
    except:
        pass
    return None

def generate_html_report(username, found_sites):
    html_content = f"""
    <html>
    <head>
        <title>Username Check Report for {username}</title>
        <style>
            body {{
                font-family: Arial, sans-serif;
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
            }}
            th, td {{
                border: 1px solid #ddd;
                padding: 8px;
                text-align: left;
            }}
            th {{
                background-color: #f2f2f2;
            }}
        </style>
    </head>
    <body>
        <h1>Username Check Report for {username}</h1>
        <table>
            <tr>
                <th>Website Name</th>
                <th>Profile URL</th>
            </tr>"""
    for site_name, uri_check in found_sites:
        html_content += f"""
            <tr>
                <td>{site_name}</td>
                <td><a href="{uri_check}" target="_blank">{uri_check}</a></td>
            </tr>"""
    html_content += """
        </table>
    </body>
    </html>"""
    with open(f"username_check_report_{username}.html", "w") as report_file:
        report_file.write(html_content)

def username_check():
    clear()
    Write.Print("[!] > Conducting Username Check...\n", default_color, interval=0)
    username = Write.Input("[?] > Enter the username: ", default_color, interval=0).strip()
    if not username:
        clear()
        Write.Print("[!] > No username provided.\n", Colors.red, interval=0)
        restart()
        return
    data = fetch_wmn_data()
    if not data:
        restart()
        return
    headers = {
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/104.0.0.0 Safari/537.36",
    }
    sites = data["sites"]
    total_sites = len(sites)
    found_sites = []
    output_accumulated = ""
    try:
        with ThreadPoolExecutor(max_workers=20) as executor:
            futures = {executor.submit(check_site, site, username, headers): site for site in sites}
            with tqdm(total=total_sites, desc="Checking sites") as pbar:
                for future in as_completed(futures):
                    try:
                        result = future.result()
                        if result:
                            site_name, uri_check = result
                            found_sites.append((site_name, uri_check))
                            found_str = f"[+] Found on: {site_name}\n[+] Profile URL: {uri_check}\n"
                            output_accumulated += found_str
                            Write.Print(found_str, Colors.green, interval=0)
                    except Exception:
                        pass
                    finally:
                        pbar.update(1)
        if found_sites:
            summary_str = f"\n[!] > Username found on {len(found_sites)} sites!\n"
            output_accumulated += summary_str
            Write.Print(summary_str, Colors.green, interval=0)
            generate_html_report(username, found_sites)
            report_msg = f"\n[!] > Report saved: username_check_report_{username}.html\n"
            output_accumulated += report_msg
            Write.Print(report_msg, Colors.green, interval=0)
        else:
            no_result_str = f"[!] > No results found for {username}.\n"
            output_accumulated += no_result_str
            Write.Print(no_result_str, Colors.red, interval=0)
    except Exception as e:
        err_str = f"[!] > An error occurred: {str(e)}\n"
        output_accumulated += err_str
        Write.Print(err_str, Colors.red, interval=0)
    log_option(output_accumulated)
    restart()

def reverse_phone_lookup(phone_number):
    query = phone_number
    url = "https://www.googleapis.com/customsearch/v1"
    params = {
        'key': API_KEY,
        'cx': CX,
        'q': query,
        'num': 5
    }
    results_data = []
    try:
        response = requests.get(url, params=params, timeout=10)
        response.raise_for_status()
        data = response.json()
        if 'items' in data and len(data['items']) > 0:
            items = data['items'][:5]
            for idx, item in enumerate(items, start=1):
                page_url = item.get('link')
                if page_url:
                    page_text = fetch_page_text(page_url)
                    results_data.append((idx, page_url, page_text))
        else:
            results_data.append((1, None, "No results found."))
    except requests.exceptions.Timeout:
        results_data.append((1, None, "Request timed out."))
    except requests.exceptions.HTTPError as e:
        results_data.append((1, None, f"HTTP error: {e.response.status_code}"))
    except Exception as e:
        results_data.append((1, None, f"Error: {str(e)}"))
    clear()
    info_text = f"""
╭─{' '*78}─╮
|{' '*28}Reverse Phone Lookup{' '*28}|
|{'='*80}|
| [+] > Phone: {phone_number:<66}|
|{'-'*80}|
"""
    for idx, link, content in results_data:
        info_text += f"| Result #{idx:<2}{' '*(73-len(str(idx)))}|\n"
        if link:
            truncated_link = (link[:75] + "...") if len(link) > 75 else link
            info_text += f"| Link: {truncated_link:<72}|\n"
        info_text += f"|{'-'*78}|\n"
        lines = [content[i:i+78] for i in range(0, len(content), 78)]
        for line in lines:
            info_text += f"| {line:<78}|\n"
        if idx != results_data[-1][0]:
            info_text += f"|{'='*78}|\n"
    info_text += f"╰─{' '*78}─╯"
    Write.Print(info_text, Colors.white, interval=0)
    log_option(info_text)
    restart()

def check_ssl_cert(domain):
    try:
        context = ssl.create_default_context()
        with socket.create_connection((domain, 443), timeout=10) as sock:
            with context.wrap_socket(sock, server_hostname=domain) as ssock:
                cert = ssock.getpeercert()
        subject = dict(x[0] for x in cert['subject'])
        issued_to = subject.get('commonName', 'N/A')
        issuer = dict(x[0] for x in cert['issuer'])
        issued_by = issuer.get('commonName', 'N/A')
        not_before = cert['notBefore']
        not_after = cert['notAfter']
        not_before_dt = datetime.strptime(not_before, "%b %d %H:%M:%S %Y %Z")
        not_after_dt = datetime.strptime(not_after, "%b %d %H:%M:%S %Y %Z")
        info_text = f"""
╭─{' '*78}─╮
|{' '*33}SSL Certificate Info{' '*32}|
|{'='*80}|
| [+] > Domain:       {domain:<58}|
| [+] > Issued To:    {issued_to:<58}|
| [+] > Issued By:    {issued_by:<58}|
| [+] > Valid From:   {str(not_before_dt):<58}|
| [+] > Valid Until:  {str(not_after_dt):<58}|
╰─{' '*78}─╯
"""
        Write.Print(info_text, Colors.white, interval=0)
        log_option(info_text)
    except ssl.SSLError as e:
        Write.Print(f"[!] > SSL Error: {str(e)}\n", Colors.red, interval=0)
    except socket.timeout:
        Write.Print("[!] > Connection timed out.\n", Colors.red, interval=0)
    except Exception as e:
        Write.Print(f"[!] > An error occurred retrieving SSL cert info: {str(e)}\n", Colors.red, interval=0)
    restart()

def check_robots_and_sitemap(domain):
    urls = [
        f"https://{domain}/robots.txt",
        f"https://{domain}/sitemap.xml"
    ]
    result_text = f"""
╭─{' '*78}─╮
|{' '*32}Site Discovery{' '*32}|
|{'='*80}|
| [+] > Domain:  {domain:<63}|
|{'-'*80}|
"""
    for resource_url in urls:
        try:
            resp = requests.get(resource_url, timeout=10)
            if resp.status_code == 200:
                lines = resp.text.split('\n')
                result_text += f"| Resource: {resource_url:<66}|\n"
                result_text += f"| Status: 200 (OK)\n"
                result_text += f"|{'-'*80}|\n"
                snippet = "\n".join(lines[:10])
                snippet_lines = snippet.split('\n')
                for sline in snippet_lines:
                    trunc = sline[:78]
                    result_text += f"| {trunc:<78}|\n"
                if len(lines) > 10:
                    result_text += "| ... (truncated)\n"
            else:
                result_text += f"| Resource: {resource_url:<66}|\n"
                result_text += f"| Status: {resp.status_code}\n"
            result_text += f"|{'='*80}|\n"
        except requests.exceptions.RequestException as e:
            result_text += f"| Resource: {resource_url}\n"
            result_text += f"| Error: {str(e)}\n"
            result_text += f"|{'='*80}|\n"
    result_text += f"╰─{' '*78}─╯"
    Write.Print(result_text, Colors.white, interval=0)
    log_option(result_text)
    restart()

def check_dnsbl(ip_address):
    dnsbl_list = [
        "zen.spamhaus.org",
        "bl.spamcop.net",
        "dnsbl.sorbs.net",
        "b.barracudacentral.org"
    ]
    reversed_ip = ".".join(ip_address.split(".")[::-1])
    results = []
    for dnsbl in dnsbl_list:
        query_domain = f"{reversed_ip}.{dnsbl}"
        try:
            answers = dns.resolver.resolve(query_domain, 'A')
            for ans in answers:
                results.append((dnsbl, str(ans)))
        except dns.resolver.NXDOMAIN:
            pass
        except dns.resolver.NoAnswer:
            pass
        except Exception as e:
            results.append((dnsbl, f"Error: {str(e)}"))
    report = f"""
╭─{' '*78}─╮
|{' '*33}DNSBL Check{' '*34}|
|{'='*80}|
| [+] > IP: {ip_address:<67}|
|{'-'*80}|
"""
    if results:
        report += "| The IP is listed on the following DNSBL(s):\n"
        for dnsbl, answer in results:
            report += f"|   {dnsbl:<25} -> {answer:<45}|\n"
    else:
        report += "| The IP is NOT listed on the tested DNSBL(s).\n"
    report += f"╰─{' '*78}─╯"
    Write.Print(report, Colors.white, interval=0)
    log_option(report)
    restart()

def fetch_webpage_metadata(url):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36)"
    }
    try:
        resp = requests.get(url, headers=headers, timeout=10)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "lxml")
        title_tag = soup.find("title")
        meta_desc = soup.find("meta", attrs={"name": "description"})
        meta_keyw = soup.find("meta", attrs={"name": "keywords"})
        title = title_tag.get_text(strip=True) if title_tag else "N/A"
        description = meta_desc["content"] if meta_desc and "content" in meta_desc.attrs else "N/A"
        keywords = meta_keyw["content"] if meta_keyw and "content" in meta_keyw.attrs else "N/A"
        result_text = f"""
╭─{' '*78}─╮
|{' '*31}Webpage Metadata{' '*31}|
|{'='*80}|
| [+] > URL:         {url:<58}|
| [+] > Title:       {title:<58}|
| [+] > Description: {description:<58}|
| [+] > Keywords:    {keywords:<58}|
╰─{' '*78}─╯
"""
        Write.Print(result_text, Colors.white, interval=0)
        log_option(result_text)
    except Exception as e:
        Write.Print(f"[!] > Error fetching metadata: {str(e)}\n", Colors.red, interval=0)
    restart()

PERPLEXITY_API_URL = "https://api.perplexity.ai/chat/completions"
PERPLEXITY_API_KEY = "INSERT PERPLEXITY API KEY HERE"
perplexity_headers = {
    "Authorization": f"Bearer {PERPLEXITY_API_KEY}",
    "Content-Type": "application/json",
}

def business_search():
    clear()
    Write.Print("[!] > Retrieve information about a business.\n", default_color, interval=0)
    business_name = Write.Input("[?] > Enter the business or persons name to search:", default_color, interval=0).strip()
    if not business_name:
        Write.Print("[!] > No business name was provided.\n", Colors.red, interval=0)
        restart()
        return
    payload_business_info = {
        "model": "llama-3.1-sonar-huge-128k-online",
        "messages": [
            {
                "role": "system",
                "content":(
                    "You are a business search assistant. Your job is to obtain and display details about a business "
                    "or business owner. You must give detailed answers nd have a well formatted,structured output                                                                                                                                                          that"
                    "explains all aspects of the search query. Provide thorough, well-structured answers."
                )
            },
            {
                "role": "user",
                "content": f"Provide me with general information about {business_name}."
            }
        ],
        "max_tokens": 50000,
        "temperature": 1.1,
    }
    out_text = ""
    try:
        response = requests.post(PERPLEXITY_API_URL, headers=perplexity_headers, json=payload_business_info)
        if response.status_code == 200:
            data = response.json()
            out_text = "\nGeneral Business Information:\n" + data["choices"][0]["message"]["content"] + "\n"
            Write.Print(out_text, Colors.white, interval=0)
        else:
            err_msg = f"Error: {response.status_code}, {response.text}\n"
            out_text = err_msg
            Write.Print(err_msg, Colors.red, interval=0)
    except Exception as e:
        out_text = f"[!] > Exception in retrieving business info: {str(e)}\n"
        Write.Print(out_text, Colors.red, interval=0)
    log_option(out_text)
    restart()

def travel_assessment(location):
    clear()
    Write.Print("[!] > Creating a comprehensive travel risk analysis based on the location you searched...Please wait...\n", default_color, interval=0)
    analysis = ""
    try:
        openai.api_key = "INSERT OPENAI API KEY HERE"
        prompt = f"""
Provide a comprehensive, highly detailed travel risk analysis for the following location: {location}.

Along with each of the following categories, include practical advice, specific examples,
local nuances, relevant data sources, disclaimers, and recommended best practices
wherever applicable. Each output should have at least 350 words. Each output is required to include the following in the output - General Details About The Risk Assessment (at least 350 characters, no more than 425), Best Practices (at least 350 characters, no more than 425), Practical Guidelines (at least 350 characters, no more than 425), Relevant Disclaimers (at least 350 characters, no more than 425)

1) Political stability and security situation
2) Crime rates and common criminal activities
3) Natural disaster risk and seasonal weather patterns
4) Health risks and medical facility standards
5) Local laws and cultural considerations
6) Transportation infrastructure quality
7) Communication infrastructure reliability
8) Current travel advisories from relevant authorities
9) Required equipment or valuable items
10) Entry requirements and visa status
11) Vaccination requirements
12) Currency restrictions
13) Insurance coverage needs
14) Accommodation security standards
15) Local transportation options
16) Communication plans and backup methods
17) Data security requirements
18) Local emergency contact numbers
19) Embassy/consulate locations and contacts
20) Medical evacuation procedures
21) Natural disaster response protocols
22) Communication tree and escalation procedures
23) Emergency funds access
24) Backup documentation storage
25) Meeting points and safe locations
26) Required medical preparations
27) Food and water safety guidelines
28) Local medical facility locations
29) Insurance coverage verification
30) Mental health support access
31) COVID-19 and infectious disease protocols
32) Environmental hazard precaution
33) Political situation updates
34) Weather and natural disaster alerts
35) Health and disease outbreak tracking
36) Transportation disruption monitoring
37) Civil unrest tracking
38) Crime pattern monitoring
39) Travel advisory updates
40) Local news monitoring

Include disclaimers about rapidly changing conditions, and note that official government websites
and reputable sources (e.g., WHO, CDC, local government portals) should be consulted for the most
up-to-date information. Make sure to answer all 40 sections and do not omit information for brevity. Your answers must be comprehensive and not one sentence long. There are no space constraints. Your answers to general details, best practices, practical guidelines, and diclaimers should be comprehensive and at least 3-5 sentences long each. This is mandatory. Answer all 40 parameters.
"""
        response = openai.ChatCompletion.create(
            model="o1-preview",
            messages=[{"role": "user", "content": prompt}],
            temperature=1
        )
        analysis = response.choices[0].message.content
        Write.Print(analysis, Colors.white, interval=0)
    except Exception as e:
        analysis = f"[!] > An error occurred: {str(e)}\n"
        Write.Print(analysis, Colors.red, interval=0)
    log_option(analysis)
    restart()

def botometer_search():
    clear()
    username = Write.Input("[?] > Enter a X/Twitter username (with or without @): ", default_color, interval=0).strip()
    if not username:
        Write.Print("[!] > No username was provided.\n", Colors.red, interval=0)
        restart()
        return
    if not username.startswith("@"):
        username = "@" + username
    Write.Print(f"[!] > Checking Botometer score for {username}...\n", default_color, interval=0)
    output_text = ""
    try:
        url = "https://botometer-pro.p.rapidapi.com/botometer-x/get_botscores_in_batch"
        payload = {
            "user_ids": [],
            "usernames": [username]
        }
        headers = {
            "x-rapidapi-key": "INSERT BOTOMETER / RAPIDAPI KEY HERE",
            "x-rapidapi-host": "botometer-pro.p.rapidapi.com",
            "Content-Type": "application/json"
        }
        response = requests.post(url, json=payload, headers=headers, timeout=10)
        result = response.json()
        output_text = json.dumps(result, indent=2)
        Write.Print(output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > An error occurred: {str(e)}"
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def hudson_rock_email_infection_check():
    clear()
    email = Write.Input("[?] > Enter an email to check infection status: ", default_color, interval=0).strip()
    if not email:
        clear()
        Write.Print("[!] > No email was provided.\n", Colors.red, interval=0)
        restart()
        return
    output_text = ""
    try:
        url = "https://cavalier.hudsonrock.com/api/json/v2/osint-tools/search-by-email"
        params = {"email": email}
        resp = requests.get(url, params=params, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        output_lines = []
        output_lines.append(f"[+] Hudson Rock email infection check results for {email}:\n")
        if isinstance(data, dict):
            for k, v in data.items():
                output_lines.append(f"{k}: {v}")
        else:
            output_lines.append("No structured data available.")
        output_text = "\n".join(output_lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except requests.exceptions.Timeout:
        output_text = "[!] > Request timed out when contacting Hudson Rock.\n"
        clear()
        Write.Print(output_text, default_color, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, default_color, interval=0)
    log_option(output_text)
    restart()

def hudson_rock_username_infection_check():
    clear()
    username = Write.Input("[?] > Enter a username to check infection status: ", default_color, interval=0).strip()
    if not username:
        clear()
        Write.Print("[!] > No username was provided.\n", Colors.red, interval=0)
        restart()
        return
    output_text = ""
    try:
        url = "https://cavalier.hudsonrock.com/api/json/v2/osint-tools/search-by-username"
        params = {"username": username}
        resp = requests.get(url, params=params, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        output_lines = []
        output_lines.append(f"[+] Hudson Rock username infection check results for {username}:\n")
        if isinstance(data, dict):
            for k, v in data.items():
                output_lines.append(f"{k}: {v}")
        else:
            output_lines.append("No structured data available.")
        output_text = "\n".join(output_lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except requests.exceptions.Timeout:
        output_text = "[!] > Request timed out when contacting Hudson Rock.\n"
        clear()
        Write.Print(output_text, default_color, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, default_color, interval=0)
    log_option(output_text)
    restart()

def hudson_rock_domain_infection_check():
    clear()
    domain = Write.Input("[?] > Enter a domain / URL to check infection status: ", default_color, interval=0).strip()
    if not domain:
        clear()
        Write.Print("[!] > No domain was provided.\n", Colors.red, interval=0)
        restart()
        return
    output_text = ""
    try:
        url = "https://cavalier.hudsonrock.com/api/json/v2/osint-tools/search-by-domain"
        params = {"domain": domain}
        resp = requests.get(url, params=params, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        output_lines = []
        output_lines.append(f"[+] Hudson Rock domain infection check results for {domain}:\n")
        if isinstance(data, dict):
            for k, v in data.items():
                output_lines.append(f"{k}: {v}")
        else:
            output_lines.append("No structured data available.")
        output_text = "\n".join(output_lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except requests.exceptions.Timeout:
        output_text = "[!] > Request timed out when contacting Hudson Rock.\n"
        clear()
        Write.Print(output_text, default_color, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, default_color, interval=0)
    log_option(output_text)
    restart()

def hudson_rock_ip_infection_check():
    clear()
    ip_address = Write.Input("[?] > Enter IP address to check infection status: ", default_color, interval=0).strip()
    if not ip_address:
        clear()
        Write.Print("[!] > No IP provided.\n", Colors.red, interval=0)
        restart()
        return
    output_text = ""
    try:
        url = "https://cavalier.hudsonrock.com/api/json/v2/osint-tools/search-by-ip"
        params = {"ip": ip_address}
        resp = requests.get(url, params=params, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        output_lines = []
        output_lines.append(f"[+] Hudson Rock IP infection check results for {ip_address}:\n")
        if isinstance(data, dict):
            for k, v in data.items():
                output_lines.append(f"{k}: {v}")
        else:
            output_lines.append("No structured data available.")
        output_text = "\n".join(output_lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except requests.exceptions.Timeout:
        output_text = "[!] > Request timed out when contacting Hudson Rock.\n"
        clear()
        Write.Print(output_text, default_color, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, default_color, interval=0)
    log_option(output_text)
    restart()

def fact_check_text():
    clear()
    Write.Print("[!] > Enter text to fact-check:\n", default_color, interval=0)
    text_to_check = Write.Input("[?] >  ", default_color, interval=0).strip()
    if not text_to_check:
        clear()
        Write.Print("[!] > No text was provided.\n", Colors.red, interval=0)
        restart()
        return
    payload_fact_check = {
        "model": "llama-3.1-sonar-huge-128k-online",
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are an advanced AI fact-checking assistant designed to evaluate "
                    "claims and statements with rigorous accuracy and methodical analysis. "
                    "Your primary goal is to help users distinguish truth from misinformation "
                    "through careful, systematic evaluation. You must be able to apply multiple "
                    "verification methods to each claim, cross reference information across reliable "
                    "sources, check for internal consistency within claims, verify dates, numbers, "
                    "and specific details, examine original context when available, identify possible "
                    "cognitive biases, recognize emotional language that may cloud judgement, check "
                    "for cherry picked data or selective presentation, consider alternative perspectives "
                    "and explanations, and flag ideological or commercial influences. You must show and "
                    "cite all sources at the end of the output and make sure they are numbered accurately."
                )
            },
            {"role": "user", "content": f"Fact-check the following text: {text_to_check}"}
        ],
        "max_tokens": 50000,
        "temperature": 1
    }
    output_text = ""
    try:
        response = requests.post(PERPLEXITY_API_URL, headers=perplexity_headers, json=payload_fact_check)
        if response.status_code == 200:
            data = response.json()
            output_text = "\nFact Checking Results:\n" + data["choices"][0]["message"]["content"] + "\n"
            Write.Print(output_text, Colors.white, interval=0)
        else:
            err_msg = f"Error: {response.status_code}, {response.text}\n"
            output_text = err_msg
            Write.Print(err_msg, Colors.red, interval=0)
    except Exception as e:
        output_text = f"[!] > Exception in fact-checking: {str(e)}\n"
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def relationship_search():
    clear()
    Write.Print("[!] > Analyze relationships between people, organizations, or businesses:\n", default_color, interval=0)
    query = Write.Input("[?] > Enter your query: ", default_color, interval=0).strip()
    if not query:
        Write.Print("[!] > No query provided.\n", Colors.red, interval=0)
        restart()
        return
    payload_relationships = {
        "model": "llama-3.1-sonar-huge-128k-online",
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are an expert investigative researcher tasked with uncovering and analyzing connections between entities,"
                    "people, organizations, not for profits, foundations, sole proprietorships, trusts, lobbyists, partnerships, LLCs,"
                    "charities, social advocacy groups, trade and professional associations, recreational clubs, fraternal societies,"
                    "employee beneficiary associations, political action committies, govermnment agencies and departments, political parties,"
                    "soverign wealth funds, diplomatic missions / embassies, holding companies, joint ventures, cooperatives, professional corporations,"
                    "s-corporations, universities and colleges, research institutions, think tanks, academic consortiums, investment funds, hedge funds,"
                    "private equity firms, venture capital firms, mutual fund holdings, stock holdings, banking institutions, credit unions, churches,"
                    "religious orders, faith-based organizations, media companies, news organizations, broadcasting networks, hospital systems, medical practices,"
                    "healthcare consortiums, international trade organizations, intergovernmental organizations and community service groups,"
                    "For each query, thoroughly analyze the subject's background, relationships, business dealings,"
                    "partnerships, investments, board memberships, charitable activities, educational history,"
                    "professional networks and anything else. If information is speculative or unverified, clearly indicate this. Consider both direct"
                    "and indirect connections, and explain the significance of each relationship within the broader context."
                    "Flag any potential red flags or areas requiring further investigation. Your analysis should be objective,"
                    "thorough, and professional in tone, avoiding speculation while highlighting substantiated connections and"
                    "their implications. You must cite sources. Structure your response in the following format: 1) Brief overview of"
                    "the subject [EACH CLAIM MUST INCLUDE AN INLINE CITATION], 2) Key relationships and connections,"
                    "categorized by type (business, personal, philanthropic, etc.) [EVERY CONNECTION MUST BE CITED],"
                    "3) Timeline of significant interactions or partnerships [CITE SPECIFIC DATES AND SOURCES], 4)"
                    "Analysis of the strength and nature of each connection [INCLUDE EVIDENCE AND CITATIONS FOR EACH ASSESSMENT],"
                    "5) Identification of any potential conflicts of interest or notable patterns [SUPPORT WITH SPECIFIC CITATIONS],"
                    "6) Detailed textual representation of the network, personal, hobbyist and business connections. When analyzing business entities, include"
                    "parent companies, subsidiaries, joint ventures, major shareholders, and key personnel. For individuals, consider"
                    "family ties, business associates, political connections, and social networks. REQUIRED CITATION FORMAT:"
                    "Use numbered inline citations [1] and provide a complete source list at the end of your response."
                    "Each citation must include: publication name, article title, author (if available), date, and URL if applicable."
                    "Any information without a citation will be considered invalid and must be removed. If a claim combines multiple"
                    "sources, use multiple citations [1][2]. For unverified or speculative information, explicitly state Unverified: and"
                    "explain why the information lacks definitive sourcing and why it remains valuable. Your analysis should maintain strict objectivity, relying solely"
                    "on verifiable sources while highlighting substantiated connections and their implications. Circumstantial facts are allowed but should not be relied upon and only included in the report if neccessary,"
                    "END EVERY RESPONSE WITH: Sources: followed by numbered citations in Chicago style format."
                )
            },
            {"role": "user", "content": query}
        ],
        "max_tokens": 50000,
        "temperature": 1.1
    }
    output_text = ""
    try:
        response = requests.post(PERPLEXITY_API_URL, headers=perplexity_headers, json=payload_relationships)
        if response.status_code == 200:
            data = response.json()
            output_text = "\nEntity Relationship Analysis Results:\n" + data["choices"][0]["message"]["content"] + "\n"
            Write.Print(output_text, Colors.white, interval=0)
        else:
            err_msg = f"Error: {response.status_code}, {response.text}\n"
            output_text = err_msg
            Write.Print(err_msg, Colors.red, interval=0)
    except Exception as e:
        output_text = f"[!] > Exception in relationship analysis: {str(e)}\n"
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def read_file_metadata(file_path):
    clear()
    Write.Print(f"\U0001F422 Checking File Data\n {file_path}", Colors.green, interval=0)
    def timeConvert(atime):
        from datetime import datetime
        dt = atime
        newtime = datetime.fromtimestamp(dt)
        return newtime.date()
    def sizeFormat(size):
        newsize = format(size/1024, ".2f")
        return newsize + " KB"
    try:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File {file_path} does not exist.")
        Dfile = os.stat(file_path)
        file_size = sizeFormat(Dfile.st_size)
        file_name = os.path.basename(file_path)
        max_length = 60
        file_creation_time = timeConvert(Dfile.st_birthtime)
        file_modification_time = timeConvert(Dfile.st_mtime)
        file_last_Access_Date = timeConvert(Dfile.st_atime)
        mime = magic.Magic(mime=True)
        file_type = mime.from_file(file_path)
        metaData_extra = []
        def get_permission_string(file_mode):
            permissions = [
                stat.S_IRUSR, stat.S_IWUSR, stat.S_IXUSR,
                stat.S_IRGRP, stat.S_IWGRP, stat.S_IXGRP,
                stat.S_IROTH, stat.S_IWOTH, stat.S_IXOTH
            ]
            labels = ['Owner', 'Group', 'Other']
            permission_descriptions = []
            for i, label in enumerate(labels):
                read = 'Yes' if file_mode & permissions[i * 3] else 'No'
                write = 'Yes' if file_mode & permissions[i * 3 + 1] else 'No'
                execute = 'Yes' if file_mode & permissions[i * 3 + 2] else 'No'
                description = f"{label} {{Read: {read}, Write: {write}, Execute: {execute}}}"
                permission_descriptions.append(description)
            return ', '.join(permission_descriptions)
        def gps_extract(exif_dict):
            gps_metadata = exif_dict['GPSInfo']
            lat_ref_num = 0
            if gps_metadata['GPSLatitudeRef'] == 'N':
                lat_ref_num += 1
            if gps_metadata['GPSLatitudeRef'] == 'S':
                lat_ref_num -= 1
            lat_list = [float(num) for num in gps_metadata['GPSLatitude']]
            lat_coordiante = (lat_list[0]+lat_list[1]/60+lat_list[2]/3600) * lat_ref_num
            long_ref_num = 0
            if gps_metadata['GPSLongitudeRef'] == 'E':
                long_ref_num += 1
            if gps_metadata['GPSLongitudeRef'] == 'W':
                long_ref_num -= 1
            long_list = [float(num) for num in gps_metadata['GPSLongitude']]
            long_coordiante = (long_list[0]+long_list[1]/60+long_list[2]/3600) * long_ref_num
            return (lat_coordiante,long_coordiante)
        permissions = get_permission_string(Dfile.st_mode)
        if(file_type.startswith("image")):
            with Image.open(file_path) as img:
                metaData_extra.append(f"|{' '*32}Image MetaData{' '*32}|")
                metaData_extra.append(f"|{'-'*78}|")
                info_dict = {
                    "Filename": img.filename,
                    "Image Size": img.size,
                    "Image Height": img.height,
                    "Image Width": img.width,
                    "Image Format": img.format,
                    "Image Mode": img.mode
                }
                for label,value in info_dict.items():
                    metaData_extra.append(f"|  {str(label):<10}: ||  {str(value)[:max_length]:<60}|")
                if img.format == 'TIFF':
                    for tag_id, value in img.tag_v2.items():
                        tag_name = TAGS.get(tag_id, tag_id)
                        metaData_extra.append(f"|  {str(tag_name):<10}: ||  {str(value)[:max_length]:<60}|")
                elif(file_path.endswith('.png')):
                    for key, value in img.info.items():
                        metaData_extra.append(f"|  {str(key):<10}: ||  {str(value)[:max_length]:<60}|")
                else:
                    imdata = img._getexif()
                    if imdata:
                        for tag_id in imdata:
                            tag = TAGS.get(tag_id, tag_id)
                            data = imdata.get(tag_id)
                            if(tag == "GPSInfo"):
                                gps = gps_extract(imdata)
                                metaData_extra.append(f"|  GPS Coordinates: ||  {gps}  |")
                                continue
                            if isinstance(data, bytes):
                                try:
                                    data = data.decode('utf-8', errors='ignore')
                                except UnicodeDecodeError:
                                    data = '<Unintelligible Data>'
                            metaData_extra.append(f"|  {str(tag):<10}: ||  {str(data)[:max_length]:<60}|")
                    else:
                        metaData_extra.append("No EXIF data found.")
        elif(file_type == "application/pdf"):
            with open(file_path, "rb") as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                pdf_data = pdf_reader.metadata
                metaData_extra.append(f"|{' '*32}PDF Metadata{' '*32}|")
                metaData_extra.append(f"|{'-'*78}|")
                if pdf_data:
                    for key, value in pdf_data.items():
                        metaData_extra.append(f"|  {str(key):<10}:  || {str(value)[:max_length]:<60}|")
                    if pdf_reader.is_encrypted:
                        metaData_extra.append(f"|  Encrypted: || Yes      |")
                    else:
                        metaData_extra.append(f"|  Encrypted: || No      |")
                else:
                    metaData_extra.append("No PDF metadata found.")
        elif(file_path.endswith(('.doc', '.docx'))):
            doc = docx.Document(file_path)
            core_properties = doc.core_properties
            doc_metadata = f"""
|{' '*32}Document Properties{' '*32}
|{'='*78}|
| Title:            || {str(core_properties.title) :<60}           |
| Author:           || {str(core_properties.author) :<60}          |
| Subject:          || {str(core_properties.subject) :<60}         |
| Keywords:         || {str(core_properties.keywords) :<60}        |
| Last Modified By: || {str(core_properties.last_modified_by) :<60}|
| Created:          || {str(core_properties.created) :<60}         |
| Modified:         || {str(core_properties.modified) :<60}        |
| Category:         || {str(core_properties.category) :<60}        |
| Content Status:   || {str(core_properties.content_status) :<60}  |
| Version:          || {str(core_properties.version) :<60}         |
| Revision:         || {str(core_properties.revision) :<60}        |
| Comments:         || {str(core_properties.comments) :<60}        |
            """
            metaData_extra.append(doc_metadata)
        elif(file_path.endswith(('.xlsx', '.xlsm'))):
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            properties = workbook.properties
            excel_metadata = f"""
|{' '*32}Excel Document Properties{' '*32}
|{'='*78}|
| Title:            || {str(properties.title) :<60}         |
| Author:           || {str(properties.creator) :<60}       |
| Keywords:         || {str(properties.keywords) :<60}      |
| Last Modified By: || {str(properties.lastModifiedBy) :<60}|
| Created:          || {str(properties.created) :<60}       |
| Modified:         || {str(properties.modified) :<60}      |
| Category:         || {str(properties.category) :<60}      |
| Description:      || {str(properties.description) :<60}   |
            """
            metaData_extra.append(excel_metadata)
        elif(file_path.endswith(('.pptx', '.pptm'))):
            try:
                presentation = Presentation(file_path)
                core_properties = presentation.core_properties
                pptx_metadata = f"""
|{' '*32}PowerPoint Document Properties{' '*31}|
|{'='*78}|
| Title:            || {str(core_properties.title) :<60}           |
| Author:           || {str(core_properties.author) :<60}          |
| Keywords:         || {str(core_properties.keywords) :<60}        |
| Last Modified By: || {str(core_properties.last_modified_by) :<60}|
| Created:          || {str(core_properties.created) :<60}         |
| Modified:         || {str(core_properties.modified) :<60}        |
| Category:         || {str(core_properties.category) :<60}        |
| Description:      || {str(core_properties.subject) :<60}         |
                """
                metaData_extra.append(pptx_metadata)
            except Exception as e:
                metaData_extra.append(f"[Error] Could not read PowerPoint metadata: {e}")
        elif(file_type.startswith("audio")):
            try:
                metaData_extra.append(f"|{' '*32}Audio MetaData{' '*32}|")
                metaData_extra.append(f"|{'-'*78}|")
                tinytim = TinyTag.get(file_path)
                if(tinytim):
                    metaData_extra.append(f"|  Title:    || {str(tinytim.title)[:max_length]:<60}      |")
                    metaData_extra.append(f"|  Artist:   || {str(tinytim.artist)[:max_length]:<60}     |")
                    metaData_extra.append(f"|  Genre:    || {str(tinytim.genre)[:max_length]:<60}      |")
                    metaData_extra.append(f"|  Album:    || {str(tinytim.album)[:max_length]:<60}      |")
                    metaData_extra.append(f"|  Year:     || {str(tinytim.year)[:max_length]:<60}       |")
                    metaData_extra.append(f"|  Composer: || {str(tinytim.composer)[:max_length]:<60}   |")
                    metaData_extra.append(f"|  A-Artist: || {str(tinytim.albumartist)[:max_length]:<60}|")
                    metaData_extra.append(f"|  Track     || {str(tinytim.track_total)[:max_length]:<60}|")
                    metaData_extra.append(f"|  Duration: || {f'{tinytim.duration:.2f} seconds':<60}    |")
                    metaData_extra.append(f"|  Bitrate:  || {str(tinytim.bitrate) + ' kbps':<60}       |")
                    metaData_extra.append(f"|  Samplrate:|| {str(tinytim.samplerate) + ' Hz':<60}      |")
                    metaData_extra.append(f"|  Channels: || {str(tinytim.channels):<60}                |")
                if(file_path.endswith('.mp3')):
                    audio = MP3(file_path, ID3=ID3)
                elif(file_path.endswith('.wav')):
                    audio = wave.open(file_path, 'rb')
                elif(file_path.endswith('.flac')):
                    audio = FLAC(file_path)
                elif(file_path.endswith('.ogg')):
                    audio = OggVorbis(file_path)
                elif(file_path.endswith(('.m4a', '.mp4'))):
                    audio = MP4(file_path)
                else:
                    audio = None
                if(audio is None):
                    metaData_extra.append(" \U0001F427 Cant Read Audio File for metadata.\n Unsupported")
                else:
                    if hasattr(audio, 'items') and audio.items():
                        for tag, value in audio.items():
                            metaData_extra.append(f"|  {str(tag):<10}: ||  {str(value)[:max_length]:<60}|")
            except Exception as e:
                metaData_extra.append(f"Error processing file: {str(e)}")
        clear()
        metadata_summary = f"""
|{' '*32}File Metadata{' '*33}|
|{'='*78}|
|  File Path:   || {file_path:<60}                  |
|  File Name:   || {file_name:<60}                  |
|  File Size:   || {file_size:<60}                  |
|  File Type:   || {file_type:<60}                  |
|  Permission:  || {permissions:<60}                |
|  Created:     || {str(file_creation_time):<60}    |
|  Modified:    || {str(file_modification_time):<60}|
|  Last Access: || {str(file_last_Access_Date):60}  |
"""
        metadata_summary += "\n".join(metaData_extra)
        metadata_summary += "\n" + "="*78 + "\n"
        Write.Print(metadata_summary, Colors.white, interval=0)
        log_option(metadata_summary)
    except Exception as e:
        err_msg = f" \u2620 Error reading file metadata: {str(e)}"
        Write.Print(err_msg, Colors.red, interval=0)
        log_option(err_msg)
    restart()

def hunter_domain_search():
    clear()
    Write.Print("[!] > Hunter.io Domain Search\n", default_color, interval=0)
    domain = Write.Input("[?] > Enter a domain to search via Hunter.io: ", default_color, interval=0).strip()
    if not domain:
        Write.Print("[!] > No domain provided.\n", Colors.red, interval=0)
        restart()
        return
    url = f"https://api.hunter.io/v2/domain-search?domain={domain}&api_key=INSERT API KEY HERE"
    output_text = ""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        lines = [f"[+] Hunter.io Domain Search results for {domain}:"]
        if isinstance(data, dict):
            for k, v in data.items():
                lines.append(f"{k}: {v}")
        else:
            lines.append("No structured domain data available.")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def hunter_email_finder():
    clear()
    Write.Print("[!] > Hunter.io Email Finder\n", default_color, interval=0)
    domain = Write.Input("[?] > Enter a domain (e.g. reddit.com): ", default_color, interval=0).strip()
    first_name = Write.Input("[?] > First Name: ", default_color, interval=0).strip()
    last_name = Write.Input("[?] > Last Name: ", default_color, interval=0).strip()
    if not domain or not first_name or not last_name:
        Write.Print("[!] > Missing domain or names.\n", Colors.red, interval=0)
        restart()
        return
    url = f"https://api.hunter.io/v2/email-finder?domain={domain}&first_name={first_name}&last_name={last_name}&api_key=INSERT API KEY HERE"
    output_text = ""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        lines = [f"[+] Hunter.io Email Finder results for {first_name} {last_name} @ {domain}:"]
        if isinstance(data, dict):
            for k, v in data.items():
                lines.append(f"{k}: {v}")
        else:
            lines.append("No structured email finder data available.")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def hunter_email_verifier():
    clear()
    Write.Print("[!] > Hunter.io Email Verification\n", default_color, interval=0)
    email = Write.Input("[?] > Enter an email to verify: ", default_color, interval=0).strip()
    if not email:
        Write.Print("[!] > No email provided.\n", Colors.red, interval=0)
        restart()
        return
    url = f"https://api.hunter.io/v2/email-verifier?email={email}&api_key=INSERT API KEY HERE"
    output_text = ""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        lines = [f"[+] Hunter.io Email Verification results for {email}:"]
        if isinstance(data, dict):
            for k, v in data.items():
                lines.append(f"{k}: {v}")
        else:
            lines.append("No structured verifier data available.")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def hunter_company_enrichment():
    clear()
    Write.Print("[!] > Hunter.io Company Enrichment\n", default_color, interval=0)
    domain = Write.Input("[?] > Enter a domain for enrichment (e.g. stripe.com): ", default_color, interval=0).strip()
    if not domain:
        Write.Print("[!] > No domain provided.\n", Colors.red, interval=0)
        restart()
        return
    url = f"https://api.hunter.io/v2/companies/find?domain={domain}&api_key=INSERT API KEY HERE"
    output_text = ""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        lines = [f"[+] Hunter.io Company Enrichment results for {domain}:"]
        if isinstance(data, dict):
            for k, v in data.items():
                lines.append(f"{k}: {v}")
        else:
            lines.append("No structured company data available.")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def hunter_person_enrichment():
    clear()
    Write.Print("[!] > Hunter.io Person Enrichment\n", default_color, interval=0)
    email = Write.Input("[?] > Enter an email for person enrichment: ", default_color, interval=0).strip()
    if not email:
        Write.Print("[!] > No email provided.\n", Colors.red, interval=0)
        restart()
        return
    url = f"https://api.hunter.io/v2/people/find?email={email}&api_key=INSERT API KEY HERE"
    output_text = ""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        lines = [f"[+] Hunter.io Person Enrichment results for {email}:"]
        if isinstance(data, dict):
            for k, v in data.items():
                lines.append(f"{k}: {v}")
        else:
            lines.append("No structured person data available.")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def hunter_combined_enrichment():
    clear()
    Write.Print("[!] > Hunter.io Combined Enrichment\n", default_color, interval=0)
    email = Write.Input("[?] > Enter an email for combined enrichment: ", default_color, interval=0).strip()
    if not email:
        Write.Print("[!] > No email provided.\n", Colors.red, interval=0)
        restart()
        return
    url = f"https://api.hunter.io/v2/combined/find?email={email}&api_key=INSERT API KEY HERE"
    output_text = ""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        data = resp.json()
        clear()
        lines = [f"[+] Hunter.io Combined Enrichment results for {email}:"]
        if isinstance(data, dict):
            for k, v in data.items():
                lines.append(f"{k}: {v}")
        else:
            lines.append("No structured combined data available.")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def castrick_email_search():
    clear()
    Write.Print("[!] > CastrickClues Email Search\n", default_color, interval=0)
    email = Write.Input("[?] > Enter an email to check via CastrickClues: ", default_color, interval=0).strip()
    if not email:
        Write.Print("[!] > No email provided.\n", Colors.red, interval=0)
        restart()
        return
    type_ = "email"
    query = email
    api_key = "INSERT CASTRICK API KEY HERE"
    headers = {"api-key": api_key}
    url = f"https://api.castrickclues.com/api/v1/search?query={query}&type={type_}"
    def tableify(obj, indent=0):
        lines = []
        prefix = " " * indent
        if isinstance(obj, dict):
            for key, value in obj.items():
                row_title = f"{prefix}{key}:"
                if isinstance(value, (dict, list)):
                    lines.append(f"| {row_title:<76}|")
                    lines.extend(tableify(value, indent + 2))
                else:
                    lines.append(format_table_row(row_title, str(value), indent))
        elif isinstance(obj, list):
            for idx, item in enumerate(obj):
                row_title = f"{prefix}[{idx}]:"
                if isinstance(item, (dict, list)):
                    lines.append(f"| {row_title:<76}|")
                    lines.extend(tableify(item, indent + 2))
                else:
                    lines.append(format_table_row(row_title, str(item), indent))
        else:
            lines.append(format_table_row(prefix.strip(), str(obj), indent))
        return lines
    def format_table_row(label, value, indent):
        row_lines = []
        max_inner_width = 78 - len(label) - 2
        words = value.split()
        current_line = ""
        label_prefix = f"{label} "
        for w in words:
            if len(current_line) + len(w) + 1 <= max_inner_width:
                if current_line:
                    current_line += " " + w
                else:
                    current_line = w
            else:
                row_lines.append(current_line)
                current_line = w
        if current_line:
            row_lines.append(current_line)
        lines_out = []
        if not row_lines:
            lines_out.append(f"| {label_prefix:<78}|")
        else:
            first_line = row_lines[0]
            lines_out.append(f"| {label_prefix + first_line:<78}|")
            for extra_line in row_lines[1:]:
                lines_out.append(f"| {' ' * (len(label_prefix))}{extra_line:<{78-len(label_prefix)}}|")
        return "\n".join(lines_out)
    try:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        data = response.json()
        clear()
        lines = []
        lines.append(f"╭─{' '*78}─╮")
        lines.append(f"|{' '*30}Castrick Email Search{' '*30}|")
        lines.append(f"|{'='*80}|")
        lines.append(f"| Email Queried: {email:<63}|")
        lines.append(f"|{'-'*80}|")
        if not isinstance(data, (dict, list)):
            data = {"data": data}
        table_lines = tableify(data)
        if not table_lines:
            lines.append("| No structured data returned from Castrick.|")
        else:
            lines.extend(table_lines)
        lines.append(f"╰─{' '*78}─╯")
        output_text = "\n".join(lines)
        Write.Print("\n" + output_text, Colors.white, interval=0)
    except Exception as e:
        output_text = f"[!] > Error: {str(e)}"
        clear()
        Write.Print(output_text, Colors.red, interval=0)
    log_option(output_text)
    restart()

def main():
    while True:
        try:
            clear()
            print("\033[1;31m██████╗    ██╗          █████╗     ████████╗    ███████╗     ██████╗     ██████╗     ██████╗     ███████╗")
            print("██╔════╝    ██║         ██╔══██╗    ╚══██╔══╝    ██╔════╝    ██╔════╝    ██╔═══██╗    ██╔══██╗    ██╔════╝")
            print("██║         ██║         ███████║       ██║       ███████╗    ██║         ██║   ██║    ██████╔╝    █████╗  ")
            print("██║         ██║         ██╔══██║       ██║       ╚════██║    ██║         ██║   ██║    ██╔═══╝     ██╔══╝  ")
            print("╚██████╗    ███████╗    ██║  ██║       ██║       ███████║    ╚██████╗    ╚██████╔╝    ██║         ███████╗")
            print(" ╚═════╝    ╚══════╝    ╚═╝  ╚═╝       ╚═╝       ╚══════╝     ╚═════╝     ╚═════╝     ╚═╝         ╚══════╝\033[0m")
            print("\033[1;34mC L A T S C O P E       I N F O       T O O L\033[0m   \033[1;31m(Version 1.07)\033[0m")
            author = "🛡️ By Joshua M Clatney (Clats97) - Ethical Pentesting Enthusiast 🛡️"
            Write.Print(author + "\n[OSINT]\nOpen Sources. Clear Conclusions\n", Colors.white, interval=0)
            menu = """╭─====─╮╭─===================─╮╭─============================================================─╮
|  №   ||      Function       ||                  Description                                 |
|======||=====================||==============================================================|
| [1]  || IP Address Search   || Retrieves IP address info                                    |
| [2]  || Deep Account Search || Retrieves profiles from various websites                     |
| [3]  || Phone Search        || Retrieves phone number info                                  |
| [4]  || DNS Record Search   || Retrieves DNS records (A, CNAME, MX, NS)                     |
| [5]  || Email MX Search     || Retrieves MX info for an email                               |
| [6]  || Person Name Search  || Retrieves person info via Perplexity AI                      |
| [7]  || Reverse DNS Search  || Retrieves PTR records for an IP address                      |
| [8]  || Email Header Search || Retrieves info from an email header                          |
| [9]  || Email Breach Search || Retreives email data breach info (HIBP)                      |
| [10] || WHOIS Search        || Retrieves domain registration data                           |
| [11] || Password Analyzer   || Retrieves password strength rating                           |
| [12] || Username Search     || Retrieves usernames from online accounts                     |
| [13] || Reverse Phone Search|| Retrieves references to a phone number (via Google)          |
| [14] || SSL Search          || Retrieves basic SSL certificate details from a URL           |
| [15] || Web Crawler Search  || Retrieves Robots.txt & Sitemap.xml URL file info             |
| [16] || DNSBL Search        || Retrieves IP DNS blacklist info                              |
| [17] || Web Metadata Search || Retrieves meta tags and more from a webpage                  |
| [18] || Travel Risk Search  || Retrieves a detailed travel risk assessment                  |
| [19] || Botometer Search    || Retrieves a Botometer score for an X/Twitter user account    |
| [20] || Business Search     || Retrieves general information about a business               |
| [21] || HR Email Search     || Retrieves infostealer email infection data (Hudson Rock)     |
| [22] || HR Username Search  || Retrieves infostealer username infection data (Hudson Rock)  |
| [23] || HR Domain Search    || Retrieves infostealer domain infection data (Hudson Rock)    |
| [24] || HR IP Search        || Retrieves infostealer IP address infection data (Hudson Rock)|
| [25] || Fact Check Search   || Retrieves analysis of input text for truthfulness            |
| [26] || Relationship Search || Retrieves & maps info between entities / people / businesses |
| [27] || File Metadata Search|| Retrieves metadata from various file types                   |
| [28] || Subdomain Search    || Retrieves subdomain info                                     |
| [29] || Domain Search       || Retrieves domain info (Hunter.io)                            |
| [30] || Email Search        || Retrieves email info (Hunter.io)                             |
| [31] || Email Verify Search || Retrieves email verification (Hunter.io)                     |
| [32] || Company Search      || Retrieves company enrichment (Hunter.io)                     |
| [33] || Person Info Search  || Retrieves person enrichment (Hunter.io)                      |
| [34] || Combined Search     || Retrieves combined enrichment (Hunter.io)                    |
| [35] || Email Search        || Retrieves in depth email info (CastrickClues)                |
| [0]  || Exit                || Exit ClatScope Info Tool                                     |
| [99] || Settings            || Customize ClatScope Info Tool (colour)                       |
|=============================================================================================|
╰─    ─╯╰─                   ─╯╰─                                                            ─╯
"""
            Write.Print(menu, Colors.white, interval=0)
            choice = Write.Input("[?] >  ", default_color, interval=0).strip()
            if choice == "1":
                clear()
                ip = Write.Input("[?] > IP-Address: ", default_color, interval=0)
                if not ip:
                    clear()
                    Write.Print("[!] > Enter an IP Address\n", default_color, interval=0)
                    continue
                ip_info(ip)
            elif choice == "2":
                clear()
                nickname = Write.Input("[?] > Username: ", default_color, interval=0)
                Write.Print("[!] > Conducting deep account search...\n", default_color, interval=0)
                if not nickname:
                    clear()
                    Write.Print("[!] > Enter the username\n", default_color, interval=0)
                    continue
                deep_account_search(nickname)
            elif choice == "3":
                clear()
                phone_number = Write.Input("[?] > Phone number: ", default_color, interval=0)
                if not phone_number:
                    clear()
                    Write.Print("[!] > Enter the phone number\n", default_color, interval=0)
                    continue
                phone_info(phone_number)
            elif choice == "4":
                clear()
                domain = Write.Input("[?] > Domain / URL: ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain / URL\n", default_color, interval=0)
                    continue
                Write.Print("[!] > Retrieving the DNS records...\n", default_color, interval=0)
                dns_lookup(domain)
            elif choice == "5":
                clear()
                email = Write.Input("[?] > Email: ", default_color, interval=0)
                if not email:
                    clear()
                    Write.Print("[!] > Enter email\n", default_color, interval=0)
                    continue
                email_lookup(email)
            elif choice == "6":
                clear()
                first_name = Write.Input("[?] > First Name: ", default_color, interval=0)
                last_name = Write.Input("[?] > Last Name: ", default_color, interval=0)
                city = Write.Input("[?] > City/Location: ", default_color, interval=0)
                if not first_name or not last_name:
                    clear()
                    Write.Print("[!] > Enter first and last name\n", default_color, interval=0)
                    continue
                Write.Print("[!] > Searching the persons name and location...\n", default_color, interval=0)
                person_search(first_name, last_name, city)
            elif choice == "7":
                clear()
                ip = Write.Input("[?] > Enter an IP Address for a Reverse DNS Search: ", default_color, interval=0)
                if not ip:
                    clear()
                    Write.Print("[!] > Enter an IP address\n", default_color, interval=0)
                    continue
                reverse_dns(ip)
            elif choice == "8":
                clear()
                Write.Print("[!] > Paste the raw email headers below as one single string (end with empty line):\n", default_color, interval=0)
                lines = []
                while True:
                    line = input()
                    if not line.strip():
                        break
                    lines.append(line)
                raw_headers = "\n".join(lines)
                if not raw_headers.strip():
                    clear()
                    Write.Print("[!] > No email header was provided.\n", default_color, interval=0)
                    continue
                analyze_email_header(raw_headers)
            elif choice == "9":
                clear()
                email = Write.Input("[?] > Enter an email address for a breach check: ", default_color, interval=0)
                if not email:
                    clear()
                    Write.Print("[!] > Enter an email address\n", default_color, interval=0)
                    continue
                haveibeenpwned_check(email)
            elif choice == "10":
                clear()
                domain = Write.Input("[?] > Enter a domain / URL for WHOIS lookup: ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain / URL\n", default_color, interval=0)
                    continue
                whois_lookup(domain)
            elif choice == "11":
                password_strength_tool()
            elif choice == "12":
                clear()
                Write.Print("[!] > Conducting Username Check...\n", default_color, interval=0)
                username_check()
            elif choice == "13":
                clear()
                phone_number = Write.Input("[?] > Enter phone number to reverse lookup: ", default_color, interval=0)
                if not phone_number:
                    clear()
                    Write.Print("[!] > Enter phone number\n", default_color, interval=0)
                    continue
                reverse_phone_lookup(phone_number)
            elif choice == "14":
                clear()
                domain = Write.Input("[?] > Enter a domain / URL for an SSL certificate verification: ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain name / URL\n", default_color, interval=0)
                    continue
                check_ssl_cert(domain)
            elif choice == "15":
                clear()
                domain = Write.Input("[?] > Enter domain to check for Robots.txt & Sitemap.xml file(s): ", default_color, interval=0)
                if not domain:
                    clear()
                    Write.Print("[!] > Enter a domain name / URL\n", default_color, interval=0)
                    continue
                check_robots_and_sitemap(domain)
            elif choice == "16":
                clear()
                ip_address = Write.Input("[?] > IP address to check DNSBL: ", default_color, interval=0)
                if not ip_address:
                    clear()
                    Write.Print("[!] > Enter an IP address to check DNSBL\n", default_color, interval=0)
                    continue
                check_dnsbl(ip_address)
            elif choice == "17":
                clear()
                url = Write.Input("[?] > URL for metadata extraction: ", default_color, interval=0)
                if not url:
                    clear()
                    Write.Print("[!] > Enter a URL to get metadata\n", default_color, interval=0)
                    continue
                fetch_webpage_metadata(url)
            elif choice == "18":
                clear()
                location = Write.Input("[?] > Enter location for a travel risk analysis: ", default_color, interval=0)
                if not location:
                    clear()
                    Write.Print("[!] > Enter the location\n", default_color, interval=0)
                    continue
                travel_assessment(location)
            elif choice == "19":
                clear()
                botometer_search()
            elif choice == "20":
                business_search()
            elif choice == "21":
                hudson_rock_email_infection_check()
            elif choice == "22":
                hudson_rock_username_infection_check()
            elif choice == "23":
                hudson_rock_domain_infection_check()
            elif choice == "24":
                hudson_rock_ip_infection_check()
            elif choice == "25":
                fact_check_text()
            elif choice == "26":
                relationship_search()
            elif choice == "27":
                file_path = Write.Input(" \U0001F989 Enter path to the file you want analyzed: ", default_color, interval=0)
                read_file_metadata(file_path)
            elif choice == "28":
                clear()
                domain = Write.Input("[?] > Enter domain for subdomain enumeration: ", default_color, interval=0)
                subdomain_enumeration(domain)
            elif choice == "29":
                hunter_domain_search()
            elif choice == "30":
                hunter_email_finder()
            elif choice == "31":
                hunter_email_verifier()
            elif choice == "32":
                hunter_company_enrichment()
            elif choice == "33":
                hunter_person_enrichment()
            elif choice == "34":
                hunter_combined_enrichment()
            elif choice == "35":
                castrick_email_search()
            elif choice == "0":
                clear()
                Write.Print("\n[!] > Exiting...", default_color, interval=0)
                exit()
            elif choice == "99":
                settings()
            else:
                clear()
                Write.Print("[!] > Invalid input.\n", default_color, interval=0)
        except KeyboardInterrupt:
            clear()
            Write.Print("[!] > Exiting on user request...\n", default_color, interval=0)
            exit()

def settings():
    while True:
        try:
            clear()
            print("\033[1;31m   ██████╗██╗        █████╗ ████████╗███████╗")
            print("   ██╔════╝██║       ██╔══██╗╚══██╔══╝██╔════╝")
            print("   ██║     ██║       ███████║   ██║   ███████╗")
            print("   ██║     ██║       ██╔══██║   ██║   ╚════██║")
            print("   ██████╗ ███████╗  ██║  ██║   ██║   ███████║")
            print("   ╚═════╝ ╚══════╝  ╚═╝  ╚═╝   ╚═╝   ╚══════╝\033[0m")
            print("\033[1;34mC       L      A       T       S       C       O       P       E\033[0m   \033[1;31m(Version 1.07)\033[0m")
            author = "🛡️ By Josh Clatney - Ethical Pentesting Enthusiast 🛡️"
            Write.Print(author + "\n[C.I.T]\nClatScope Info Tool\n", Colors.white, interval=0)
            settings_menu = """╭─    ─╮╭─                   ─╮╭─                                         ─╮
|  №   ||       Setting       ||                Description                |
|======||=====================||===========================================
| [1]  || Theme change        || Customize the theme
|------||---------------------||-------------------------------------------
| [0]  || Back to menu        || Exit the settings
╰─    ─╯╰─                   ─╯╰─                                         ─╯
"""
            Write.Print(settings_menu, Colors.white, interval=0)
            settings_choice = Write.Input("[?] >  ", default_color, interval=0).strip()
            if settings_choice == "1":
                change_color()
            elif settings_choice == "0":
                return
            else:
                clear()
                Write.Print("[!] > Invalid input.\n", default_color, interval=0)
        except KeyboardInterrupt:
            clear()
            Write.Print("[!] > Exiting on user request...\n", default_color, interval=0)
            exit()

if __name__ == "__main__":
    main()